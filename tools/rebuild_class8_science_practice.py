"""Build a complete Science book with supplied chapter worksheets."""
from __future__ import annotations

from collections import Counter
import hashlib
import json
import os
from pathlib import Path
import posixpath
import re
import subprocess
import sys
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from io import BytesIO

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "stage0"))
sys.path.insert(0, str(ROOT / "stage2"))

import chapter_sources
import scrape_chapters
from docx import Document
from docx.oxml.ns import qn
import pymupdf
from PIL import Image
from convert import pick_scale
from inline_content import text_runs

WORKBOOK = ROOT / "Science-notes-links.xlsx"
QUESTIONS_DIR = Path(os.environ.get('SCIENCE_INPUT_ROOT', ROOT)) / "Class 8th Science"
CHAPTER_NUMBERS = (3, 4, 7, 8, 9, 10, 12)
MODULE = "Final-Science_Class8th"


def workbook_rows(path: Path, sheet_name: str):
    """Read the small source map from OOXML, avoiding a missing openpyxl install."""
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkg": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
        selected = next((item for item in workbook.find("main:sheets", ns)
                         if item.attrib["name"] == sheet_name), None)
        if selected is None:
            raise ValueError(f"Workbook sheet not found: {sheet_name}")
        target = targets[selected.attrib[f"{{{ns['rel']}}}id"]]
        worksheet_path = (target.lstrip("/") if target.startswith("/")
                          else posixpath.normpath(posixpath.join("xl", target)))
        shared_strings = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            shared_strings = ["".join(node.text or "" for node in item.iter(f"{{{ns['main']}}}t"))
                              for item in shared_root]
        worksheet = ET.fromstring(archive.read(worksheet_path))

        def cell_value(cell):
            kind = cell.attrib.get("t")
            if kind == "inlineStr":
                return "".join(node.text or "" for node in cell.iter(f"{{{ns['main']}}}t"))
            value = cell.find("main:v", ns)
            if value is None or value.text is None:
                return ""
            if kind == "s":
                return shared_strings[int(value.text)]
            return value.text

        rows = []
        for row in worksheet.findall(".//main:sheetData/main:row", ns):
            cells = {}
            for cell in row.findall("main:c", ns):
                col = re.match(r"[A-Z]+", cell.attrib["r"]).group()
                cells[col] = cell_value(cell)
            rows.append(cells)
        if not rows:
            raise ValueError(f"Workbook sheet is empty: {sheet_name}")
        headers = tuple(rows[0].get(col, "").strip() for col in "ABCDE")
        if headers[:2] in {
            ("Chapter", "Chapter Links"),
            ("Chapters", "Science Links"),
        } and not any(headers[2:]):
            chapters = []
            for row_number, row in enumerate(rows[1:], start=2):
                label = row.get("A", "").strip()
                link = row.get("B", "").strip()
                if not label and not link:
                    continue
                if not label or not link:
                    raise ValueError(f"Workbook row {row_number} has a missing chapter or link")
                match = re.search(r"chapter\s*[- ]?\s*(\d+)", label, re.I)
                if not match:
                    raise ValueError(f"Workbook row {row_number} has no chapter number: {label}")
                number = int(match.group(1))
                url = chapter_sources.normalize_edurev_url(link, "/t/")
                slug_match = re.search(r"/t/\d+/([^/#?]+)", url)
                title = re.sub(r"^Chapter-Notes-", "", slug_match.group(1), flags=re.I) if slug_match else ""
                title = title.replace("-", " ").strip()
                chapters.append({"num": number, "name": title or label,
                                 "url": url, "formula_url": url, "infographic_url": None})
            if not chapters:
                raise ValueError(f"Workbook sheet has no chapter rows: {sheet_name}")
            return chapters
        expected = chapter_sources.WORKBOOK_HEADERS
        duplicate_name = headers == ("Chapter Name",) + expected
        if headers[:4] != expected and not duplicate_name:
            raise ValueError(f"Unexpected Science workbook headers: {headers!r}")
        chapters = []
        for row_number, row in enumerate(rows[1:], start=2):
            values = [row.get(col, "").strip() for col in "ABCD"]
            if duplicate_name:
                values = [row.get(col, "").strip() for col in "BCDE"]
                if row.get("A", "").strip():
                    values[0] = row["A"].strip()
            if not any(values):
                continue
            if any(not value for value in values[:3]):
                raise ValueError(f"Workbook row {row_number} has a missing required value")
            name, label, formula, infographic = values
            match = re.fullmatch(r"Chapter-(\d+)", label, re.I)
            if not match:
                raise ValueError(f"Invalid chapter number in workbook row {row_number}: {label}")
            number = int(match.group(1))
            if name.casefold() == "periimeter and area":
                name = "Perimeter and Area"
            chapters.append({
                "num": number,
                "name": name,
                "url": chapter_sources.normalize_edurev_url(formula, "/t/"),
                "formula_url": chapter_sources.normalize_edurev_url(formula, "/t/"),
                "infographic_url": chapter_sources.normalize_edurev_url(infographic, "/p/") if infographic else None,
            })
        if not chapters:
            raise ValueError(f"Workbook sheet has no chapter rows: {sheet_name}")
        return chapters


original_parse_sources = chapter_sources.parse_chapter_sources


def parse_sources(path, sheet=None):
    if Path(path).suffix.lower() == ".xlsx":
        return workbook_rows(Path(path), sheet)
    return original_parse_sources(path, sheet)


chapter_sources.parse_chapter_sources = parse_sources


def clean(text: str) -> str:
    text = text.replace("\u00ad", "").replace("\u2006", " ")
    text = re.sub(r"\s*-\s+", "-", text)
    return re.sub(r"[ \t]+", " ", text).strip()


def without_mark_annotation(text: str) -> str:
    """Remove a trailing mark value embedded in a worksheet prompt."""
    return re.sub(r"\s*\(\d+\)\s*$", "", clean(text))


def body(text: str, bold=False, page_break_before=False):
    element = {"type": "body", "segments": text_runs(clean(text), bool(bold))}
    if page_break_before:
        element["page_break_before"] = True
    return element


def heading(text: str, level="subsection"):
    return {"type": "heading", "level": level, "text": clean(text)}


def compact(value: str) -> str:
    value = value.translate({1: "fl", 2: "fi", 3: "fl"})
    value = value.replace("ﬂ", "fl").replace("ﬁ", "fi").replace("ﬀ", "ff")
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def fuzzy_present(target: str, source: str) -> bool:
    if target in source:
        return True
    search_at = 0
    while True:
        start = source.find(target[:25], search_at)
        if start < 0:
            return False
        cursor = start
        for char in target:
            cursor = source.find(char, cursor)
            if cursor < 0:
                break
            cursor += 1
        if cursor >= 0 and cursor - start <= len(target) * 2 + 250:
            return True
        search_at = start + 1


def question_paths(chapter: int) -> tuple[Path, Path]:
    stem = f"Chapter-{chapter}-Questions-Class8th-Science"
    return QUESTIONS_DIR / f"{stem}.docx", QUESTIONS_DIR / f"{stem}.pdf"


def row_question(row):
    number = int(row.cells[0].text.strip())
    lines = [clean(line) for paragraph in row.cells[1].paragraphs
             for line in worksheet_paragraph_text(paragraph).splitlines() if clean(line)]
    marks = row.cells[2].text.strip()
    if not lines:
        raise ValueError(f"Question {number} has no text")
    return number, lines, marks


def worksheet_paragraph_text(paragraph):
    """Retain supplied simple OMML in document order; reject unknown math."""
    math_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
    superscripts = str.maketrans('0123456789+-=()o', '⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ᵒ')
    subscripts = str.maketrans('0123456789+-=()', '₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎')

    def math_text(node):
        tag = node.tag.rsplit('}', 1)[-1]
        if tag in {'rPr', 'sSupPr', 'sSubPr', 'ctrlPr'}:
            return ''
        if tag == 't':
            return node.text or ''
        if tag in {'sSup', 'sSub'}:
            is_super = tag == 'sSup'
            script_kind = 'superscript' if is_super else 'subscript'
            base = node.find('{' + math_ns + '}e')
            exponent = node.find('{' + math_ns + '}' + ('sup' if is_super else 'sub'))
            if base is None or exponent is None:
                raise ValueError('Incomplete worksheet OMML ' + script_kind)
            power = math_text(exponent)
            allowed = '0123456789+-=()' + ('o' if is_super else '')
            if not power or any(char not in allowed for char in power):
                raise ValueError('Unsupported worksheet OMML ' + script_kind + ': ' + power)
            mapping = superscripts if is_super else subscripts
            return math_text(base).replace('\u200b', '') + power.translate(mapping)
        if tag in {'oMath', 'oMathPara', 'e', 'sup', 'sub', 'r'}:
            return ''.join(math_text(child) for child in node)
        raise ValueError('Unsupported worksheet OMML element: ' + tag)

    def content(node):
        if node.tag.startswith('{' + math_ns + '}'):
            return math_text(node)
        if node.tag == qn('w:t'):
            return node.text or ''
        if node.tag == qn('w:tab'):
            return '\t'
        if node.tag in {qn('w:br'), qn('w:cr')}:
            return '\n'
        if node.tag in {qn('w:pPr'), qn('w:rPr'), qn('w:drawing'), qn('w:pict')}:
            return ''
        return ''.join(content(child) for child in node)

    return content(paragraph._element)


def labeled_options(lines, question_number):
    options = []
    current = None
    question_lines = []
    for line in lines:
        match = re.match(r"^([a-d])[.)]\s*(.*)$", line, re.I)
        if match:
            if current is not None:
                options.append(current)
            current = match.group(2).strip()
        elif current is None:
            question_lines.append(line)
        else:
            current += " " + line
    if current is not None:
        options.append(current.strip())
    if len(options) != 4:
        raise ValueError(
            f"Question {question_number} must have four supplied options; found {len(options)}"
        )
    return question_lines, options


def source_pdf_audit(document: Document, pdf_path: Path) -> int:
    with pymupdf.open(pdf_path) as reader:
        pdf_text = compact("\n".join(page.get_text() for page in reader))
    checks = []
    for table in document.tables:
        for node in table._tbl.iter(qn("w:t")):
            value = node.text or ""
            target = compact(value)
            if len(target) >= 8:
                checks.append((value, target))
    missing = [value for value, target in checks if not fuzzy_present(target, pdf_text)]
    # Full math-bearing paragraphs must match exactly after layout normalization;
    # run-only prose checks cannot detect a lost exponent or arrow.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph._element.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'):
                        value = worksheet_paragraph_text(paragraph)
                        if compact(value) not in pdf_text:
                            missing.append(value)
                        checks.append((value, compact(value)))
    if missing:
        raise ValueError(
            f"{pdf_path.name}: {len(missing)} DOCX text fragments are absent from the supplied PDF: "
            f"{missing[:8]}"
        )
    return len(checks)


def paragraph_images(document: Document, paragraph, media_dir: Path, chapter: int):
    images = []
    extent = paragraph._element.find(f".//{qn('wp:extent')}")
    try:
        cx_emu = int(extent.get("cx", "0")) if extent is not None else 0
    except (TypeError, ValueError):
        cx_emu = 0
    base_scale = pick_scale(cx_emu)
    for blip in paragraph._element.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid or rid not in document.part.related_parts:
            raise ValueError(f"Chapter {chapter}: unresolved question image relationship {rid!r}")
        part = document.part.related_parts[rid]
        filename = f"ch{chapter}_questions_{rid}.png"
        target = media_dir / filename
        media_dir.mkdir(parents=True, exist_ok=True)
        original_filename = 'originals/' + f'ch{chapter}_questions_{rid}' + Path(str(part.partname)).suffix
        original = media_dir / original_filename
        original.parent.mkdir(parents=True, exist_ok=True)
        if original.exists() and original.read_bytes() != part.blob:
            raise ValueError('Original question media collision: ' + original_filename)
        original.write_bytes(part.blob)
        with Image.open(BytesIO(part.blob)) as opened:
            rendered = opened.convert("RGB")
            if target.exists():
                with Image.open(target) as existing:
                    if existing.convert("RGB").tobytes() != rendered.tobytes():
                        raise ValueError(f"Question image collision: {filename}")
            else:
                rendered.save(target, format="PNG", optimize=True)
        images.append({
            "type": "image",
            "filename": filename,
            "scale": base_scale * 0.5,
            "base_scale": base_scale,
            "sha256": hashlib.sha256(part.blob).hexdigest(),
            "original_filename": original_filename,
        })
    return images


def row_events(document: Document, row, media_dir: Path, chapter: int):
    events = []
    for paragraph in row.cells[1].paragraphs:
        for line in worksheet_paragraph_text(paragraph).splitlines():
            line = clean(line)
            if line:
                events.append(("text", line))
        for image in paragraph_images(document, paragraph, media_dir, chapter):
            events.append(("image", image))
    return events


def classify_question(lines, marks, number):
    first = lines[0].casefold()
    if first.startswith("state whether the given statement is true or false"):
        return "True or False"
    if first.startswith("fill in the blanks"):
        return "Fill in the Blanks"
    option_lines = [line for line in lines if re.match(r"^[a-d][.)]\s*", line, re.I)]
    if option_lines:
        if first.startswith("assertion (a):"):
            return "Assertion and Reason"
        return "Multiple Choice Questions"
    if first.startswith("match"):
        return "Match the Following"
    match = re.search(r"\d+", marks)
    if not match:
        raise ValueError(f"Question {number} has no numeric mark value")
    return {
        1: "Very Short Questions",
        2: "Short Questions",
        3: "Long Questions",
        4: "Case-Based Questions",
        5: "Very Long Questions",
    }.get(int(match.group()), "")


SOURCE_ISSUES = [{
    "chapter": 4,
    "kind": "internal_chapter_title_mismatch",
    "source": "Chapter-4-Questions-Class8th-Science.docx/pdf",
    "detail": "The internal title says Chapter 3; filename, workbook mapping, and electricity content identify Chapter 4.",
    "resolution": "Placed under Chapter 4 with the user's explicit confirmation; source files were not edited.",
}, {
    "chapter": 3,
    "kind": "duplicate_generic_prompt",
    "source_questions": [31, 32],
    "detail": "Both supplied prompts are exactly 'Match the following:' and have distinct source figures.",
    "resolution": "Both were preserved in source order.",
}]

original_download_images = scrape_chapters.download_content_images
original_check_structure = scrape_chapters.check_structure
questions_inserted = set()
expected_practice_text = {}
question_metadata = {}
question_images = []


def parse_case_events(events):
    output = []
    parts = []
    removed_marks = []
    omitted_or_count = 0
    for kind, value in events:
        if kind != 'text':
            if parts:
                raise ValueError('Case figure appears between subquestions; review source order')
            output.append((kind, value))
            continue
        if value.strip().casefold() == 'or':
            omitted_or_count += 1
            continue
        match = re.search(r'\s*\(\s*(\d+)\s*\)\s*$', value)
        if match:
            marks = int(match.group(1))
            text = value[:match.start()].strip()
            removed_marks.append({'source': value, 'marks': marks, 'display': text})
            parts.append({'text': text, 'marks': marks, 'answer_lines': marks})
        elif parts:
            raise ValueError('Unmarked case content follows marked parts: ' + value)
        else:
            output.append((kind, value))
    if not parts:
        text_positions = [index for index, event in enumerate(output) if event[0] == 'text']
        boundary = next((index for index in text_positions if 'answer the following:' in output[index][1].casefold()), None)
        if boundary is None:
            raise ValueError('Cannot identify case subquestions')
        passage = output[:boundary + 1]
        for kind, value in output[boundary + 1:]:
            if kind != 'text':
                raise ValueError('Unmarked case figure after question boundary')
            parts.append({'text': value, 'marks': None, 'answer_lines': 1})
        output = passage
    output.append(('subquestions', parts))
    return output, {'removed_marks': removed_marks, 'omitted_or_count': omitted_or_count,
                    'case_part_count': len(parts)}


def practice_elements(chapter: int, media_dir: Path):
    from science_table_manifest import convert_table
    docx_path, pdf_path = question_paths(chapter)
    document = Document(docx_path)
    if len(document.tables) < 2:
        raise ValueError(f'{docx_path.name} does not contain its question table')
    pdf_checks = source_pdf_audit(document, pdf_path)
    rows = [row for row in document.tables[1].rows if row.cells[0].text.strip().isdigit()]
    if [int(row.cells[0].text.strip()) for row in rows] != list(range(1, len(rows) + 1)):
        raise ValueError('Question rows are not contiguous')
    class_number = 6 if '6th' in QUESTIONS_DIR.name else 8
    result = [dict(type='heading', level='section', text='Practice Questions', page_break_before=True)]
    expected = []
    type_counts = Counter()
    rendered_counts = Counter()
    current_heading = None
    local_images = []
    tables = []
    cases = []
    answer_line_count = 0
    for row in rows:
        number, lines, marks = row_question(row)
        title = classify_question(lines, marks, number)
        if not title:
            raise ValueError('Unsupported question type')
        if title != current_heading:
            result.append(heading(title))
            expected.append(title)
            current_heading = title
        type_counts[title] += 1
        raw_events = row_events(document, row, media_dir, chapter)
        events = []
        for kind, value in raw_events:
            if kind == 'image':
                converted = convert_table(value, class_number, chapter)
                if converted['type'] == 'table':
                    tables.append(converted)
                    events.append(('table', converted))
                else:
                    local_images.append(value)
                    events.append((kind, value))
            else:
                events.append((kind, value))
        if title in {'Multiple Choice Questions', 'Assertion and Reason'}:
            prompt_lines, options = labeled_options(lines, number)
            if title == 'Assertion and Reason' and (len(prompt_lines) < 2 or not prompt_lines[1].casefold().startswith('reason (r):')):
                raise ValueError(f'Chapter {chapter} question {number}: incomplete assertion/reason pair')
            content = [body(' '.join(prompt_lines))]
            content.extend(value for kind, value in events if kind in {'image', 'table'})
            content.append({'type': 'options', 'values': options})
            result.append({'type': 'question', 'content': content})
            expected.extend(prompt_lines + options)
            rendered_counts[title] += 1
            continue
        if title in {'True or False', 'Fill in the Blanks'}:
            result.append(body(lines[0]))
            expected.append(lines[0])
            for line in lines[1:]:
                content = [body(line)]
                if title == 'Fill in the Blanks':
                    content[0]['ragged_right'] = True
                if title == 'True or False':
                    content.append({'type': 'answer_lines', 'count': 1})
                    answer_line_count += 1
                result.append({'type': 'question', 'content': content})
                expected.append(line)
                rendered_counts[title] += 1
            continue
        if title == 'Case-Based Questions':
            events, audit = parse_case_events(events)
            cases.append({'source_question': number, **audit})
        elif len(events) > 2:
            last_asset = max((index for index, (kind, value) in enumerate(events) if kind != 'text'), default=0)
            trailing = events[last_asset + 1:]
            if len(trailing) > 1 and all(kind == 'text' and value.endswith('?') for kind, value in trailing):
                parts = [{'text': value, 'marks': None, 'answer_lines': 1} for kind, value in trailing]
                events = events[:last_asset + 1] + [('subquestions', parts)]
        content = []
        for kind, value in events:
            if kind == 'text':
                content.append(body(value))
                expected.append(value)
            elif kind == 'subquestions':
                items = []
                for part in value:
                    expected.append(part['text'])
                    count = part['answer_lines']
                    items.append({'content': [body(part['text']), {'type': 'answer_lines', 'count': count}]})
                    answer_line_count += count
                content.append({'type': 'subquestions', 'items': items})
            else:
                content.append(value)
        if title != 'Case-Based Questions' and not any(kind == 'subquestions' for kind, value in events):
            count = 1 if title in {'Very Short Questions', 'Short Questions'} else 3 if title in {'Long Questions', 'Very Long Questions'} else 0
            if count:
                content.append({'type': 'answer_lines', 'count': count})
                answer_line_count += count
        result.append({'type': 'question', 'content': content})
        rendered_counts[title] += 1
    for table in tables:
        expected.extend(cell for row in table['rows'] for cell in row if cell and not cell.isdigit())
    expected_table_counts = {6: {3: 2, 4: 1, 6: 2, 7: 3, 11: 2, 12: 2}, 8: {3: 4, 4: 1}}
    if len(tables) != expected_table_counts[class_number].get(chapter, 0):
        raise ValueError(f'Chapter {chapter}: reviewed table inventory changed')
    question_images.extend({'chapter': chapter, **image} for image in local_images)
    metadata = dict(source_docx=docx_path.name, source_pdf=pdf_path.name,
                    source_rows=len(rows), rendered_items=sum(rendered_counts.values()),
                    counts_by_type=dict(type_counts), rendered_items_by_type=dict(rendered_counts),
                    source_figures=len(local_images), docx_pdf_prose_checks=pdf_checks,
                    converted_tables=tables, cases=cases, answer_line_count=answer_line_count)
    return result, expected, metadata


def add_practice_after_chapter_images(elements, chapter, session, media, args, report):
    original_download_images(elements, chapter, session, media, args, report)
    number = chapter["num"]
    if number in CHAPTER_NUMBERS:
        if number in questions_inserted:
            raise ValueError(f"Chapter {number} practice questions were inserted twice")
        questions, expected, metadata = practice_elements(number, media)
        elements.extend(questions)
        expected_practice_text[number] = expected
        question_metadata[number] = metadata
        questions_inserted.add(number)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def tracked_sha256(relative_path: Path) -> str:
    """Hash a file as stored in HEAD without relying on a sibling clean clone."""
    data = subprocess.check_output(
        ["git", "-c", f"safe.directory={ROOT.as_posix()}", "-C", str(ROOT),
         "show", f"HEAD:{relative_path.as_posix()}"],
    )
    return hashlib.sha256(data).hexdigest()


def source_hashes():
    hashes = {WORKBOOK.name: sha256(WORKBOOK)}
    if tracked_sha256(Path(WORKBOOK.name)) != hashes[WORKBOOK.name]:
        raise ValueError("Workbook does not match the checked-out Git revision")
    for chapter in CHAPTER_NUMBERS:
        for path in question_paths(chapter):
            hashes[path.name] = sha256(path)
            relative = Path("Class 8th Science") / path.name
            if tracked_sha256(relative) != hashes[path.name]:
                raise ValueError(f"Source does not match the checked-out Git revision: {path.name}")
    return hashes


def verify_practice_before_publish(chapters, directory, module, outputs):
    result = original_check_structure(chapters, directory, module, outputs)
    if questions_inserted != set(CHAPTER_NUMBERS):
        raise ValueError(f"Practice questions were not inserted for all chapters: {questions_inserted}")
    import pymupdf
    from validate_against_pdf import map_pdf, norm

    rendered_by_output = {}
    for label, path in outputs.items():
        with pymupdf.open(path) as pdf:
            ranges = map_pdf(pdf, chapters, generated=True)
            for chapter, page_range in zip(chapters, ranges):
                chapter_pages = list(range(page_range['start_page'] - 1, page_range['end_page']))
                numeric_words = Counter(word[4] for page_number in chapter_pages
                                        for word in pdf[page_number].get_text('words'))
                expected_numeric = Counter(cell for table in question_metadata[chapter['num']]['converted_tables']
                                           for row in table['rows'] for cell in row if cell.isdigit())
                for cell, count in expected_numeric.items():
                    if numeric_words[cell] < count:
                        raise ValueError(f'{label}: table value {cell!r} appears {numeric_words[cell]} times; expected at least {count}')
                placed = [info['digest'] for page_number in range(page_range['start_page'] - 1, page_range['end_page'])
                          for info in pdf[page_number].get_image_info(hashes=True)]
                cursor = 0
                for image in question_images:
                    if image['chapter'] != chapter['num']:
                        continue
                    digest = pymupdf.Pixmap(directory / 'media' / image['filename']).digest
                    try:
                        cursor = placed.index(digest, cursor) + 1
                    except ValueError as error:
                        raise ValueError(f'{label}: lost or reordered question figure {image["filename"]}') from error
            rendered_by_output[label] = {
                chapter["num"]: scrape_chapters.rendered_text(
                    pdf,
                    page_range,
                    [chapter["name"], "Practice Questions"],
                )
                for chapter, page_range in zip(chapters, ranges)
            }

    for chapter in CHAPTER_NUMBERS:
        checks = {}
        for label, rendered_by_chapter in rendered_by_output.items():
            expected = expected_practice_text[chapter]
            rendered = rendered_by_chapter[chapter]
            missing = [text for text in expected if norm(text) not in rendered]
            if missing:
                raise ValueError(
                    f"{label}: Chapter {chapter} Practice Questions text missing from compiled PDF: "
                    f"{missing[:8]}"
                )
            checks[label] = {"source_text_items_checked": len(expected), "missing": 0}
        question_metadata[chapter]["compiled_prose_checks"] = checks

    tex = (directory / f"{module}.tex").read_text(encoding="utf-8")
    lyx = (directory / f"{module}.lyx").read_text(encoding="utf-8")
    expected_lines = sum(item['answer_line_count'] for item in question_metadata.values())
    if tex.count(r'\rule{\linewidth}{0.3pt}') != expected_lines:
        raise ValueError('TeX answer-line count differs from audited count')
    if lyx.count('rule{\\backslash\nlinewidth}{0.3pt}') != expected_lines:
        raise ValueError('LyX answer-line count differs from audited count')
    for metadata in question_metadata.values():
        for table in metadata['converted_tables']:
            from question_layout import tex_content, lyx_content
            import convert
            if tex_content([table], convert.render_tex_segments, convert.tex_escape) not in tex:
                raise ValueError('TeX table cells or ordering changed')
            if lyx_content([table], convert.render_lyx_segments, convert.tex_escape) not in lyx:
                raise ValueError('LyX native table cells or ordering changed')
            reference = 'media/' + table['source_image']['filename']
            if reference in tex or reference in lyx:
                raise ValueError('Converted table picture is still rendered: ' + reference)
            with Image.open(directory / reference) as original:
                original.verify()
            original_path = directory / 'media' / table['source_image']['original_filename']
            if sha256(original_path) != table['source_image']['sha256']:
                raise ValueError('Original table media hash differs from reviewed source')
    for image in question_images:
        if abs(image['scale'] / image['base_scale'] - .5) > 1e-9:
            raise ValueError('Question figure sizing factor changed')
        expected_tex = f"width={image['scale']}\\textwidth]{{media/{image['filename']}}}"
        expected_lyx = f"filename media/{image['filename']}"
        expected_lyx_width = f"width {int(round(image['scale'] * 100))}text%"
        if expected_tex not in tex or expected_lyx not in lyx or expected_lyx_width not in lyx:
            raise ValueError(f"Question image scale/reference mismatch: {image['filename']}")

    result["practice_questions"] = {
        "chapters": {str(ch): question_metadata[ch] for ch in CHAPTER_NUMBERS},
        "total_source_rows": sum(item["source_rows"] for item in question_metadata.values()),
        "total_rendered_items": sum(item["rendered_items"] for item in question_metadata.values()),
        "question_images": {"count": len(question_images), "scale_factor": 0.5},
        "source_issues": SOURCE_ISSUES,
        "source_hashes": source_hashes(),
        "converted_table_count": sum(len(item['converted_tables']) for item in question_metadata.values()),
        "case_part_count": sum(case['case_part_count'] for item in question_metadata.values() for case in item['cases']),
        "omitted_or_count": sum(case['omitted_or_count'] for item in question_metadata.values() for case in item['cases']),
        "answer_line_count": sum(item['answer_line_count'] for item in question_metadata.values()),
    }
    return result


scrape_chapters.download_content_images = add_practice_after_chapter_images
scrape_chapters.check_structure = verify_practice_before_publish


def git_commit() -> str:
    return subprocess.check_output(
        ["git", "-c", f"safe.directory={ROOT.as_posix()}", "-C", str(ROOT),
         "rev-parse", "HEAD"],
        text=True,
    ).strip()


def update_manifest():
    path = Path(os.environ.get('SCIENCE_OUTPUT_DIR', ROOT / "stage2" / "output")) / MODULE / "manifest.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    manifest.update({
        "question_image_scale": 0.5,
        "question_sources": {
            "git_commit": git_commit(),
            "git_dirty": True,
            "source_hashes": source_hashes(),
            "adapter_sha256": sha256(Path(__file__)),
            "question_layout_sha256": sha256(ROOT / 'stage2' / 'question_layout.py'),
            "table_manifest_sha256": sha256(ROOT / 'tools' / 'science_table_manifest.py'),
            "convert_sha256": sha256(ROOT / "stage2" / "convert.py"),
            "scrape_chapters_sha256": sha256(ROOT / "stage0" / "scrape_chapters.py"),
        },
    })
    path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


if __name__ == "__main__":
    args = [
        "--links", str(WORKBOOK),
        "--sheet", "Class 8th",
        "--module", MODULE,
        "--title", "Class 8th Science",
        "--output-dir", os.environ.get('SCIENCE_OUTPUT_DIR', str(ROOT / "stage2" / "output")),
        "--cache-dir", os.environ.get('SCIENCE_CACHE_DIR', str(ROOT / "input" / "scraped")),
        "--offline",
    ]
    status = scrape_chapters.main(args)
    if status == 0:
        update_manifest()
    raise SystemExit(status)
