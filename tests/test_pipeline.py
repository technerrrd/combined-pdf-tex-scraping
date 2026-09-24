import io
from pathlib import Path
from unittest.mock import Mock
import pytest
import requests
from PIL import Image
import build_support as support
import inline_content as inline
from html_source import fit_image_scale, parse_html
from scrape_chapters import arguments, module_name, prose_chunks, rendered_text, select_chapters
from release_science import CLASSES, workbook_rows
from validate_against_pdf import coverage, map_pdf, reference_lines, validate_ranges, norm
from chapter_sources import (compare_pdf_pages, extract_infographic_pdf_url,
                             find_reference_pdf, parse_chapter_sources)


def response(status=200, content=b'good', headers=None):
    r = Mock(status_code=status, content=content, headers=headers or {})
    if status >= 400: r.raise_for_status.side_effect = requests.HTTPError(str(status))
    return r


def test_cache_validated_atomic_and_offline(tmp_path):
    session = Mock(); session.get.return_value = response()
    path = support.cache_path(tmp_path, 'https://a/ch1', 'pages')
    validate = lambda data: None if data == b'good' else (_ for _ in ()).throw(ValueError('bad'))
    assert support.cached_fetch(session, 'https://a/ch1', path, validate) == b'good'
    session.get.reset_mock()
    assert support.cached_fetch(session, 'https://a/ch1', path, validate, offline=True) == b'good'
    session.get.assert_not_called()
    path.write_bytes(b'bad')
    with pytest.raises(support.BuildError, match='invalid cached'): support.cached_fetch(session, 'url', path, validate, offline=True)
    assert support.cache_path(tmp_path, 'https://b/ch1', 'pages') != path


def test_retries_and_refresh(tmp_path, monkeypatch):
    sleeps = []; monkeypatch.setattr(support.time, 'sleep', sleeps.append)
    session = Mock(); session.get.side_effect = [response(429, headers={'Retry-After':'3'}), requests.Timeout(), response()]
    path = tmp_path/'page'; path.write_bytes(b'old')
    assert support.cached_fetch(session, 'url', path, lambda d: None, refresh=True) == b'good'
    assert sleeps == [3, 2]
    assert session.get.call_count == 3


@pytest.mark.parametrize('status', [401,403,404,500])
def test_failed_download_does_not_replace_cache(tmp_path, monkeypatch, status):
    monkeypatch.setattr(support.time, 'sleep', lambda _: None)
    session=Mock(); session.get.return_value=response(status)
    path=tmp_path/'page'; path.write_bytes(b'old')
    with pytest.raises(support.BuildError): support.cached_fetch(session,'url',path,lambda _:None,refresh=True)
    assert path.read_bytes() == b'old'
    assert session.get.call_count == (3 if status == 500 else 1)


def test_atomic_write_failure_preserves_old(tmp_path,monkeypatch):
    path=tmp_path/'cache'; path.write_bytes(b'old')
    monkeypatch.setattr(support.os,'replace',Mock(side_effect=OSError('interrupted')))
    with pytest.raises(OSError): support.atomic_write(path,b'new')
    assert path.read_bytes()==b'old'
    assert list(tmp_path.iterdir())==[path]


def test_offline_missing_and_corrupt_image(tmp_path):
    with pytest.raises(support.BuildError,match='offline cache miss'): support.cached_fetch(Mock(),'url',tmp_path/'none',lambda _:None,offline=True)
    with pytest.raises(OSError): support.validate_image(b'html error')
    data=io.BytesIO(); Image.new('RGB',(10,10)).save(data,format='PNG'); support.validate_image(data.getvalue())


@pytest.mark.parametrize('html',['<title>Access denied</title><p>paywall</p>','<h2>Empty</h2>','<input type=password><p>Login</p>'])
def test_blocked_and_empty(html):
    with pytest.raises(ValueError): parse_html(html)


def test_html_order_spacing_tables_and_lists():
    els=parse_html('<article><h2>1. Heading</h2><p><strong>active</strong> and plain<p>Before<img src="https://x/a_lg.jpg">After<ul><li>Parent<ol><li>Nested</ol><li>Last</ul><table><tr><td>CO<sub>2</sub></td></tr></table></article>')
    assert els[0]['text']=='Heading'
    assert els[1]['text']=='active and plain'
    assert [e['type'] for e in els[2:5]]==['body','image','body']
    lists=[e for e in els if e['type']=='list']
    assert [(e['ordered'],e['items'][0][0]) for e in lists]==[(False,0),(True,1),(False,0)]
    assert els[-1]['rows'][0][0][-1]['kind']=='sub'


def test_sparse_legacy_images_use_article_content_and_skip_promotions():
    html = '''<div class="contenttextdiv"><table><tr><td>Clear all your doubts with EduRev</td></tr></table>
    <div class="contenttextdiv"><h2>1. Real topic</h2><p>Substantive notes.</p>
    <img src="https://cn.edurev.in/ApplicationImages/Temp/abc_sp.png"></div></div>'''
    els = parse_html(html)
    assert [e['type'] for e in els] == ['heading', 'body', 'image']
    assert els[-1]['filename'] == 'abc_sp.png'


def test_number_prefixed_content_images_are_kept_but_decorations_are_skipped():
    html = '''<article><h2>Topic</h2><p>Substantive notes.</p>
    <img alt="Thermometer" src="https://cn.edurev.in/ApplicationImages/Temp/1421561_75514512-4eb3-45e9-9b2b-656447055db2_lg.png">
    <img alt="Decoration" src="https://cn.edurev.in/cdn_lib/Vector.png"></article>'''
    images = [element for element in parse_html(html) if element['type'] == 'image']
    assert len(images) == 1
    assert images[0]['filename'].startswith('1421561_')


def test_instructional_gif_is_accepted_for_static_frame_conversion():
    html = '''<article><h2>Nutrition</h2><p>Substantive notes.</p>
    <img alt="Nutrition" src="https://cn.edurev.in/ApplicationImages/Temp/27602a0a-ab1d-4a9b-8222-62c0e350892f_lg.gif"></article>'''
    images = [element for element in parse_html(html) if element['type'] == 'image']
    assert len(images) == 1
    assert images[0]['filename'].endswith('_lg.gif')


def test_promo_table_is_skipped_before_span_validation():
    html = '<article><table class="coursedatatable"><tr><td colspan="2">Science for Class 6 91 videos | 431 docs</td></tr></table><h2>Topic</h2><p>Notes.</p></article>'
    assert [e['type'] for e in parse_html(html)] == ['heading', 'body']


@pytest.mark.parametrize('source,expected', [('CO₂','textsubscript{2}'),('x²','textsuperscript{2}'),('α → β',r'\alpha'),('√(a²+b²)',r'\surd'),(r'\(A\rightarrow B\)',r'\rightarrow'),(r'\(\frac{a}{b}\)',r'\frac{a}{b}'),(r'\[\sqrt{x}\]',r'\sqrt{x}'),('Price $5 and 50%','\\$5')])
def test_notation(source,expected):
    runs=inline.text_runs(source)
    assert expected in inline.tex(runs)
    if any(s['kind']=='math' for s in runs): assert '\\begin_inset Formula' in inline.lyx(runs)


def test_mathml_and_annotation():
    html='<p><math><mfrac><mi>a</mi><msqrt><mi>b</mi></msqrt></mfrac></math></p>'
    assert r'\frac{a}{\sqrt{b}}' in inline.tex(parse_html(html)[0]['segments'])
    html='<p><span class="katex"><span>duplicate</span><math><semantics><mi>x</mi><annotation encoding="application/x-tex">x^2</annotation></semantics></math></span></p>'
    runs=parse_html(html)[0]['segments']; assert len(runs)==1 and runs[0]['text']=='x^2'


def test_plain_text_inside_mathml_remains_explicit_math():
    runs = parse_html('<p>For example, <math>9 - 5 = 4</math>.</p>')[0]['segments']
    assert any(run['kind'] == 'math' and run['text'] == '9 - 5 = 4' for run in runs)
    with pytest.raises(ValueError, match='Empty explicit equation'):
        parse_html('<p><math></math></p>')


@pytest.mark.parametrize('formula',[r'\input{secret}',r'\write18{bad}',r'\begin{document}x\end{document}','{x',r'\newcommand{x}{y}'])
def test_unsupported_math(formula):
    with pytest.raises(ValueError): inline.check_math(formula)


def test_unknown_mathml():
    with pytest.raises(ValueError,match='Unsupported MathML'): parse_html('<p><math><maction><mi>x</mi></maction></math></p>')


def test_full_line_coverage_and_boundary():
    prefix='A sufficiently long sentence with a common beginning '
    result=coverage([prefix+'correct ending'],prefix+'wrong ending',1,'tex',{})
    assert result['coverage']==0
    lines=[f'Unique statement number {i} ends here' for i in range(20)]
    assert coverage(lines,' '.join(lines[:19]),1,'tex',{})['coverage']==.95
    assert norm('CO₂  ﬁne')==norm('CO2 fine')
    with pytest.raises(support.ValidationError): coverage([], '',1,'tex',{})


def test_mapping_never_guesses():
    doc=Mock(); doc.get_toc.return_value=[[1,'Other',1]]; doc.__len__=Mock(return_value=10)
    with pytest.raises(support.ValidationError,match='Ambiguous'): map_pdf(doc,[{'num':1,'name':'Missing'}])
    chapters=[{'num':1,'name':'One'},{'num':2,'name':'Two'}]
    ranges=[dict(chapter=1,start_page=1,end_page=2),dict(chapter=2,start_page=3,end_page=5)]
    assert validate_ranges(ranges,chapters,5)==ranges
    ranges[1]['start_page']=2
    with pytest.raises(support.ValidationError): validate_ranges(ranges,chapters,5)


def test_cli_selection_and_module():
    chapters=[dict(num=i) for i in [3,1,2]]
    assert [c['num'] for c in select_chapters(chapters,'1,3')]==[3,1]
    with pytest.raises(ValueError): select_chapters(chapters,'4')
    with pytest.raises(ValueError): select_chapters([dict(num=1),dict(num=1)],None)
    with pytest.raises(ValueError): module_name(arguments(['--module','../bad']))
    with pytest.raises(ValueError): module_name(arguments(['old','--module','new']))
    assert module_name(arguments(['--links','MyNotes.txt']))=='MyNotes'


def test_xlsx_chapter_sources_normalize_and_preserve_order(tmp_path):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Class 6th'
    sheet.append(['Chapter Name', 'Chapter-Number',
                  'Important Points and Formulas link', 'Infographic link'])
    sheet.append(['Periimeter and Area', 'Chapter-6',
                  'https://edurev.in/t/1/formulas', 'https://edurev.in/p/2/infographic'])
    sheet.append([None, None, None, None])
    sheet.append(['Fractions', 'Chapter-7',
                  'edurev.in/t/3/fractions', 'https://edurev.in/p/4/fractions'])
    path = tmp_path / 'links.xlsx'
    workbook.save(path)
    chapters = parse_chapter_sources(path, 'Class 6th')
    assert [chapter['num'] for chapter in chapters] == [6, 7]
    assert chapters[0]['name'] == 'Perimeter and Area'
    assert chapters[1]['formula_url'].startswith('https://edurev.in/t/3/')
    assert chapters[0]['infographic_url'].startswith('https://edurev.in/p/2/')


def test_xlsx_chapter_sources_reject_invalid_schema(tmp_path):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.title = 'Class 6th'
    workbook.active.append(['Wrong', 'Headers'])
    path = tmp_path / 'bad.xlsx'
    workbook.save(path)
    with pytest.raises(ValueError, match='headers'):
        parse_chapter_sources(path, 'Class 6th')
    with pytest.raises(ValueError, match='--sheet is required'):
        parse_chapter_sources(path)


@pytest.mark.parametrize('duplicate_name', [False, True])
def test_xlsx_formula_only_chapters_and_duplicate_display_name(tmp_path, duplicate_name):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Maths'
    headers = ['Chapter Name', 'Chapter-Number',
               'Important Points and Formulas link', 'Infographic link']
    row = ['Connecting the Dots', 'Chapter-13', 'https://edurev.in/t/13/dots', None]
    if duplicate_name:
        headers.insert(0, 'Chapter Name')
        row.insert(0, 'Connecting the Dots...')
    sheet.append(headers)
    sheet.append(row)
    path = tmp_path / 'formulas.xlsx'
    workbook.save(path)
    chapters = parse_chapter_sources(path, 'Maths')
    assert len(chapters) == 1
    assert chapters[0]['num'] == 13
    assert chapters[0]['name'] == ('Connecting the Dots...' if duplicate_name else 'Connecting the Dots')
    assert chapters[0]['infographic_url'] is None
    # A formula URL remains required even when no infographic is supplied.
    sheet.cell(2, 4 if duplicate_name else 3).value = None
    workbook.save(path)
    with pytest.raises(ValueError, match='missing required value'):
        parse_chapter_sources(path, 'Maths')


def test_infographic_pdf_extraction_is_exact_and_host_limited():
    url = 'https://cn.edurev.in/files/1421561_example.pdf'
    assert extract_infographic_pdf_url(f'<script>viewer = "{url}"</script>') == url
    with pytest.raises(ValueError, match='exactly one'):
        extract_infographic_pdf_url('<p>No PDF</p>')
    with pytest.raises(ValueError, match='exactly one'):
        extract_infographic_pdf_url(f'{url} https://cn.edurev.in/files/other.pdf')


def test_reference_pdf_resolution_and_visual_comparison(tmp_path):
    import pymupdf
    formula = tmp_path / 'Chapter-6-Important-formulas.pdf'
    infographic = tmp_path / 'Infographics-Chapter-6.pdf'
    with pymupdf.open() as document:
        page = document.new_page(width=300, height=500)
        page.draw_rect((20, 20, 280, 480), color=(1, 0, 0), width=3)
        page.insert_text((40, 60), 'Perimeter and Area')
        document.save(infographic)
    formula.write_bytes(infographic.read_bytes())
    assert find_reference_pdf(tmp_path, 6, 'formula') == formula
    assert find_reference_pdf(tmp_path, 6, 'infographic') == infographic
    assert compare_pdf_pages(infographic.read_bytes(), infographic) == [0.0]


def test_infographic_writers_use_one_page_uncropped_layout(tmp_path):
    import convert
    media = tmp_path / 'media'
    media.mkdir()
    elements = [
        dict(type='heading', level='section', text='Infographics', page_break_before=True),
        dict(type='infographic', filename='page1.png', page_number=1, page_count=2),
        dict(type='infographic', filename='page2.png', page_number=2, page_count=2),
    ]
    tex_path = tmp_path / 'book.tex'
    lyx_path = tmp_path / 'book.lyx'
    convert.write_tex(elements, tex_path, media)
    convert.write_lyx(elements, lyx_path, media)
    tex = tex_path.read_text()
    lyx = lyx_path.read_text()
    assert tex.index(r'\section{Infographics}') < tex.index('media/page1.png') < tex.index('media/page2.png')
    assert 'height=.80\\textheight,keepaspectratio' in tex
    assert lyx.index('Infographics') < lyx.index('media/page1.png') < lyx.index('media/page2.png')
    assert '\theight 70page%\n\tkeepAspectRatio' in lyx


def test_body_page_break_is_rendered_by_both_writers(tmp_path):
    import convert
    media = tmp_path / 'media'
    media.mkdir()
    elements = [
        dict(type='body', segments=[('Case Study 2', True)], page_break_before=True),
    ]
    tex_path = tmp_path / 'book.tex'
    lyx_path = tmp_path / 'book.lyx'
    convert.write_tex(elements, tex_path, media)
    convert.write_lyx(elements, lyx_path, media)
    tex = tex_path.read_text()
    lyx = lyx_path.read_text()
    assert tex.index('\\clearpage') < tex.index('Case Study 2')
    assert lyx.index('clearpage') < lyx.index('Case Study 2')


def test_mcq_writers_preserve_math_in_prompt_and_options(tmp_path):
    import convert
    media = tmp_path / 'media'
    media.mkdir()
    elements = [dict(
        type='mcq', text='ignored', options=['ignored'] * 4,
        question_segments=[inline.run('Which fraction equals ', kind='text'),
                           inline.run(r'\frac{3}{5}', kind='math')],
        option_segments=[[inline.run(r'\frac{3}{5}', kind='math')],
                         [inline.run('50%', kind='text')],
                         [inline.run('75%', kind='text')],
                         [inline.run('60%', kind='text')]],
    ), dict(type='body', segments=[inline.run('Question text '),
                                    inline.run(r'\frac{3}{5}', kind='math'),
                                    inline.run(' continues after the fraction.')])]
    tex_path = tmp_path / 'book.tex'
    lyx_path = tmp_path / 'book.lyx'
    convert.write_tex(elements, tex_path, media)
    convert.write_lyx(elements, lyx_path, media)
    tex = tex_path.read_text()
    lyx = lyx_path.read_text()
    assert r'Which fraction equals $\frac{3}{5}$' in tex
    assert r'A) $\frac{3}{5}$' in tex
    assert r'Question text ' + '\n' + r'\begin_inset Formula $\frac{3}{5}$' in lyx
    assert r'\backslash' + '\nfrac{3}{5}' in lyx
    assert r'\backslash' + '\n%' in lyx


def test_publication_rollback(tmp_path,monkeypatch):
    old=tmp_path/'module'; old.mkdir(); (old/'good').write_text('old')
    staged=tmp_path/'staged'; staged.mkdir()
    original=Path.rename
    def fail(self,target):
        if self==staged: raise OSError('publication interrupted')
        return original(self,target)
    monkeypatch.setattr(Path,'rename',fail)
    with pytest.raises(OSError): support.publish(staged,old)
    assert (old/'good').read_text()=='old'

def test_compiler_timeout_retains_diagnostics(tmp_path,monkeypatch):
    import subprocess
    monkeypatch.setattr(support.subprocess,'run',Mock(side_effect=subprocess.TimeoutExpired('compiler',1,output=b'partial')))
    with pytest.raises(support.BuildError,match='timed out'): support.command(['compiler'],tmp_path,tmp_path/'compiler.log',timeout=1)
    assert b'partial' in (tmp_path/'compiler.log').read_bytes()


def test_compile_documents_uses_lyx_export_switch(tmp_path, monkeypatch):
    seen = []

    def fake_command(args, directory, log, timeout=180, environment=None):
        seen.append(list(args))
        if args[0] == 'lyx':
            if len([item for item in seen if item[0] == 'lyx']) == 1:
                userdir = Path(args[args.index('-userdir') + 1])
                userdir.mkdir(parents=True, exist_ok=True)
                (userdir / 'lyxrc.defaults').write_text('configured')
                raise support.BuildError('first-run configuration')
            assert '-E' in args
            assert '-e' not in args
            (directory / args[-2]).write_text('\\documentclass{book}\\begin{document}\\end{document}')
            return ''
        job = next(arg.split('=', 1)[1] for arg in args if arg.startswith('-jobname='))
        (directory / f'{job}.log').write_text('')
        import pymupdf
        pdf = pymupdf.open()
        pdf.new_page()
        pdf.save(directory / f'{job}.pdf')
        pdf.close()
        return ''

    monkeypatch.setattr(support, 'command', fake_command)
    outputs = support.compile_documents(tmp_path, 'book', {'lyx': 'lyx', 'pdflatex': 'pdflatex'})
    assert seen[0][seen[0].index('-batch') + 1:seen[0].index('-batch') + 3] == ['-E', 'pdflatex']
    assert len([item for item in seen if item[0] == 'lyx']) == 2
    assert set(outputs) == {'tex', 'lyx'}


def test_compiler_missing_is_not_skipped(monkeypatch):
    monkeypatch.setattr(support.shutil,'which',lambda _:None)
    with pytest.raises(support.BuildError,match='Missing compiler'): support.executable('nonexistent-compiler')


def test_headless_linux_compiler_environment():
    linux = support.compiler_environment({'PATH': '/usr/bin'}, platform_name='posix')
    assert linux['QT_QPA_PLATFORM'] == 'offscreen'
    assert linux['openin_any'] == linux['openout_any'] == 'p'
    desktop = support.compiler_environment({'DISPLAY': ':0'}, platform_name='posix')
    assert 'QT_QPA_PLATFORM' not in desktop
    explicit = support.compiler_environment({'QT_QPA_PLATFORM': 'minimal'}, platform_name='posix')
    assert explicit['QT_QPA_PLATFORM'] == 'minimal'


def test_compiler_python_environment_adds_windows_py_shim(tmp_path):
    environment, shim = support.compiler_python_environment(
        tmp_path, {'PATH': 'C:\\Windows'}, platform_name='nt', python_executable='C:\\Python312\\python.exe')
    assert environment['PYTHON'].endswith('Python312\\python.exe')
    assert environment['PATH'].split(support.os.pathsep)[0] == str(shim)
    assert 'Python312' in (shim / 'py.cmd').read_text()


def test_latex_diagnostics_report_only_actionable_layout_problems():
    log = '''Overfull \\hbox (4.0pt too wide) in paragraph
Overfull \\hbox (12.5pt too wide) in paragraph
LaTeX Warning: Float too large for page by 20.0pt.
Missing character: There is no ₹ in font cmr10!'''
    diagnostics = support.latex_diagnostics(log, 'tex')
    assert [item['kind'] for item in diagnostics] == ['overfull_hbox', 'oversized_float', 'missing_glyph']
    assert [item['severity'] for item in diagnostics] == ['warning', 'error', 'error']


def test_unsupported_and_matrix_mathml():
    for markup,expected in [('<mroot><mi>x</mi><mn>3</mn></mroot>',r'\sqrt[3]{x}'),('<msubsup><mi>x</mi><mn>1</mn><mn>2</mn></msubsup>','{x}_{1}^{2}'),('<mfenced><mi>x</mi><mi>y</mi></mfenced>',r'\left(x,y\right)'),('<mover><mi>x</mi><mo>^</mo></mover>',r'\hat{x}')]:
        runs=parse_html('<p><math>'+markup+'</math></p>')[0]['segments']
        assert expected in inline.tex(runs)
    assert norm('a ≤ b') != norm('a ≥ b')
    assert norm(r'a \leq b') == norm('a ≤ b')


def test_mathjax_script_and_literal_escaping():
    elements=parse_html('<p>Equation <script type="math/tex">x^2</script> here.</p>')
    assert r'$x^2$' in inline.tex(elements[0]['segments'])
    assert inline.escape('\\{50%}')==r'\textbackslash{}\{50\%\}'
    assert inline.plain(parse_html('<p>Visible<!-- hidden --> text.</p>')[0]['segments'])=='Visible text.'


def test_math_closer_split_across_bold_boundary_is_rejoined():
    runs = parse_html(r'<p>A fixed <strong>ratio \(2:1\</strong>) is used.</p>')[0]['segments']
    assert inline.plain(runs) == r'A fixed ratio 2:1 is used.'
    assert any(run['kind'] == 'math' and run['text'] == '2:1' for run in runs)


def test_standard_right_arrow_command_is_preserved():
    assert inline.check_math(r'A \rightarrow B') == r'A \rightarrow B'


def test_docx_fallback_legacy_segments(tmp_path):
    from docx import Document
    import convert
    document=Document(); document.add_paragraph('A plain paragraph with '); document.paragraphs[0].add_run('bold text').bold=True
    source=tmp_path/'source.docx'; document.save(source)
    elements,*_=convert.parse_docx(source)
    media=tmp_path/'media'; media.mkdir()
    convert.write_tex(elements,tmp_path/'fallback.tex',media)
    convert.write_lyx(elements,tmp_path/'fallback.lyx',media)
    assert r'\textbf{bold text}' in (tmp_path/'fallback.tex').read_text()
    assert '\\series bold' in (tmp_path/'fallback.lyx').read_text()
    assert r'\raggedbottom' in (tmp_path/'fallback.tex').read_text()
    assert r'\raggedbottom' in (tmp_path/'fallback.lyx').read_text()


def test_original_chapter_numbers_compact_spacing_and_unicode_rupee(tmp_path):
    import convert
    media = tmp_path/'media'; media.mkdir()
    elements = [dict(type='heading', level='chapter', text='Seven', number=7),
                dict(type='body', text='₹ 7', segments=[dict(kind='text', text='₹ 7')])]
    tex_path = tmp_path/'numbered.tex'; lyx_path = tmp_path/'numbered.lyx'
    convert.write_tex(elements, tex_path, media)
    convert.write_lyx(elements, lyx_path, media)
    tex = tex_path.read_text(encoding='utf-8'); lyx = lyx_path.read_text(encoding='utf-8')
    assert r'\setcounter{chapter}{6}' in tex
    assert r'\newunicodechar{₹}{\rupee}' in tex
    assert r'\titlespacing*{\section}{0pt}{1.25ex' in tex
    assert '\\begin_layout Standard\n\\begin_inset ERT' in lyx
    assert 'setcounter{chapter}{6}' in lyx
    assert r'\newunicodechar{₹}{\rupee}' in lyx
    assert '₹ 7' in tex and '₹ 7' in lyx


def test_lyx_chapter_banner_uses_title_size_that_fits_long_science_names(tmp_path):
    import convert
    convert.copy_template_assets(convert.TEMPLATE_DIR, tmp_path)
    structure = (tmp_path / 'structure.tex').read_text(encoding='utf-8')
    assert r'node {\Large\sffamily\bfseries\color{black}\thechapter. #1\strut}' in structure
    assert r'node {\huge\sffamily\bfseries\color{black}\thechapter. #1\strut}' not in structure


@pytest.mark.parametrize(('requested','width','height','expected'), [
    (.75, 1200, 400, .39),   # landscape: width cap, then 25% reduction
    (.60, 600, 600, .346),   # square: height cap, then 25% reduction
    (.60, 300, 600, .173),   # portrait: height cap, then 25% reduction
    (.60, 150, 900, .058),   # extremely tall: height cap, then 25% reduction
    (.25, 1200, 400, .188),  # smaller source-requested size is also reduced
])
def test_balanced_image_scale(requested, width, height, expected):
    assert fit_image_scale(requested, width, height) == expected


def test_image_scale_can_be_overridden_and_is_validated():
    assert fit_image_scale(.60, 600, 600, 1) == .462
    with pytest.raises(ValueError, match='scale factor'): fit_image_scale(.60, 600, 600, 0)
    assert arguments([]).image_scale == .75


def test_publication_keeps_previous_snapshot(tmp_path):
    destination=tmp_path/'module';destination.mkdir();(destination/'file').write_text('old')
    staged=tmp_path/'staged';staged.mkdir();(staged/'file').write_text('new')
    support.publish(staged,destination)
    assert (destination/'file').read_text()=='new'
    assert next(tmp_path.glob('module.previous-*')).joinpath('file').read_text()=='old'


def test_duplicate_pdf_markers_and_scanned_reference(tmp_path):
    import pymupdf
    from validate_against_pdf import validate_reference
    path=tmp_path/'empty.pdf'
    with pymupdf.open() as doc:
        doc.new_page();doc.set_toc([[1,'Chapter Notes: One',1],[1,'Chapter Notes: One',1]])
        assert map_pdf(doc,[dict(num=1,name='One')])==[dict(chapter=1,start_page=1,end_page=1)]
        doc.save(path)
    with pytest.raises(support.ValidationError,match='no extractable'):
        validate_reference(path,[dict(num=1,name='One',elements=[])],{})

def test_images_inside_formatted_lists_are_not_lost():
    elements=parse_html('<article><p>Introduction.</p><ul><li><strong>Before <span><img src="https://x/test_lg.jpg"></span> after</strong></li></ul></article>')
    assert [e['type'] for e in elements]==['body','list','image','list']
    assert elements[1]['items'][0][1][0]['bold']
    assert elements[3]['items'][0][1][0]['bold']


def test_table_spans_fail_instead_of_flattening():
    with pytest.raises(ValueError,match='Complex table'):
        parse_html('<table><tr><td colspan="2">Merged cell</td></tr></table>')


def test_compiler_failure_does_not_publish(tmp_path,monkeypatch):
    import scrape_chapters
    from build_support import atomic_write,cache_path
    url='https://example.org/t/1/Chapter-Notes-One'
    links=tmp_path/'links';links.write_text(url+' - Chapter1')
    output=tmp_path/'out';good=output/'notes';good.mkdir(parents=True);(good/'original').write_text('keep')
    cache=tmp_path/'cache';atomic_write(cache_path(cache,url,'pages'),b'<p>A complete sample paragraph.</p>')
    monkeypatch.setattr(support,'preflight',lambda: {})
    monkeypatch.setattr(support,'compile_documents',Mock(side_effect=support.BuildError('compiler failure')))
    args=['--links',str(links),'--module','notes','--cache-dir',str(cache),'--output-dir',str(output),'--offline']
    assert scrape_chapters.main(args)==1
    assert (good/'original').read_text()=='keep'
    assert any('compiler failure' in p.read_text() for p in output.glob('notes-*/report.txt'))


def test_html_invalid_fetch_not_cached(tmp_path):
    session=Mock();session.get.return_value=response(content=b'<title>Access denied</title><p>Blocked</p>')
    path=tmp_path/'cache'
    with pytest.raises(support.BuildError):support.cached_fetch(session,'url',path,lambda data:parse_html(data.decode()))
    assert not path.exists()


def test_html_repairs_math_delimiter_split_by_formatting():
    html = '<main><h2>Compounds</h2><p>A <strong>ratio \\(2:1\\</strong>) is fixed.</p></main>'
    elements = parse_html(html, 'https://example.test/chapter')
    paragraph = next(item for item in elements if item['type'] == 'body')
    assert inline.plain(paragraph['segments']) == 'A ratio 2:1 is fixed.'
    assert any(run['kind'] == 'math' and run['text'] == '2:1' for run in paragraph['segments'])


def test_retry_after_http_date():
    from datetime import datetime,timezone,timedelta
    from email.utils import format_datetime
    delay=support.retry_delay(format_datetime(datetime.now(timezone.utc)+timedelta(seconds=20)),0)
    assert 18 <= delay <= 20

def test_legacy_currency_rendering():
    assert inline.tex([('Rs. 250',False)])==r'\rupee~ 250'
    assert inline.lyx([('₹250',False)]) == '₹250'

def test_two_currency_amounts_are_not_inferred_as_equations():
    runs=inline.text_runs('Prices are $5 and $10.')
    assert all(s['kind']=='text' for s in runs)
    assert inline.tex(runs)==r'Prices are \$5 and \$10.'

def test_reference_counts_the_generated_chapter_title(tmp_path):
    import pymupdf
    from validate_against_pdf import validate_reference
    source=tmp_path/'reference.pdf'
    with pymupdf.open() as doc:
        page=doc.new_page();page.insert_text((72,72),'Unique Heading\nA substantive body sentence.')
        doc.set_toc([[1,'Unique Heading',1]]);doc.save(source)
    chapters=[dict(num=1,name='Unique Heading',elements=[dict(type='body',text='A substantive body sentence.',segments=[('A substantive body sentence.',False)])])]
    assert validate_reference(source,chapters,{})[0]['coverage']==1

def test_edurev_navigation_table_is_skipped_before_span_validation():
    result=parse_html('<article><table class="tbl_cntnt"><tr><td>Table of Contents</td></tr><tr><td colspan="2">Navigation</td></tr></table><h2>Chapter topic</h2><p>Substantive chapter content.</p></article>')
    assert [e['type'] for e in result]==['heading','body']

def test_notes_exclude_embedded_quiz_widgets():
    elements=parse_html('<article><p>Real notes.</p><div id="content_questions"><p>Question and solution widget</p></div><p>More notes.</p></article>')
    assert [e['text'] for e in elements]==['Real notes.','More notes.']


def test_notes_exclude_optional_fact_callouts_and_all_quiz_variants():
    html = '''<article>
      <p>Real notes before.</p>
      <blockquote><p><strong>Fun fact!</strong></p><p>Optional thermometer fact.</p></blockquote>
      <div class="question_block_42"><p>Try yourself: A question?</p><p>Option A</p></div>
      <blockquote><strong>Do you know?</strong><p>Optional seed fact.</p></blockquote>
      <p>The teacher asked, "Do you know the answer?" This is ordinary prose.</p>
    </article>'''
    elements = parse_html(html)
    assert [e['text'] for e in elements] == [
        'Real notes before.',
        'The teacher asked, "Do you know the answer?" This is ordinary prose.',
    ]


def test_typographic_dashes_match_tex_punctuation():
    assert norm('Health—not just disease')==norm('Health---not just disease')
    assert norm('books-almost') == norm('books-\nalmost')
    assert norm('well-being') == norm('well being')
    assert norm('12-4') != norm('12-5')
    assert norm('20 °C') == norm('20 ◦C')
    assert norm('Water boils at 100 °C') == norm('Water boils at 100 C')
    assert norm(r'Water boils at 100 \circ C') == norm('Water boils at 100 °C')
    assert norm('Saptaṛiṣhi, Dhruva tārā, Sūrya') == norm('Saptar.is.hi Dhruva t�ar�a S�urya')
    assert norm('Saptaṛiṣhi and Sūrya') == norm('Saptar.is.hi and S¯urya')
    assert norm('Re\x01ectional \x02gure on the \x03oor') == norm('Reflectional figure on the floor')


def test_release_workbook_preserves_class_sheet_order():
    workbook = Path(__file__).parents[1] / 'Science-notes-links.xlsx'
    rows = workbook_rows(workbook, [CLASSES[number]['sheet'] for number in (6, 7, 8)])
    assert [number for number, _ in rows['Class 6th']] == [7, 8, 9, 10, 11, 12]
    assert [number for number, _ in rows['Class 7th']] == [4, 5, 6, 7, 8, 9, 10]
    assert [number for number, _ in rows['Class 8th']] == [3, 4, 7, 8, 9, 10, 12]


def test_reference_coverage_ignores_multilevel_heading_numbers_and_corrupt_logo():
    result = coverage(['1.2 Formula (Regular Polygon)', 'UREV', 'of 6'],
                      'Formula (Regular Polygon)', 9, 'parsed', {})
    assert result['coverage'] == 1


def test_reference_lines_exclude_embedded_quiz():
    page = Mock()
    page.rect.height = 800
    def line(text, y): return {'bbox': (0, y, 100, y + 10), 'spans': [{'text': text, 'size': 10, 'font': 'Arial'}]}
    page.get_text.return_value = {'blocks': [{'lines': [
        line('Notes before quiz.', 100), line('MULTIPLE CHOICE QUESTION', 120),
        line('Try yourself: A question?', 140), line('An option.', 160),
        line('View Solution', 180), line('Notes after quiz.', 200),
    ]}]}
    assert reference_lines(page) == ['Notes before quiz.', 'Notes after quiz.']


def test_rendered_text_ignores_page_furniture_inside_a_paragraph():
    class Rect:
        height = 800

    class Page:
        rect = Rect()
        def __init__(self, lines): self.lines = lines
        def get_text(self, kind):
            assert kind == 'dict'
            return {'blocks': [{'lines': [
                {'bbox': (0, y, 100, y + 10), 'spans': [{'text': text}]}
                for y, text in self.lines
            ]}]}

    doc = [
        Page([(700, 'Temperature often affects how much solute a solvent can dissolve.')]),
        Page([(40, '10'), (60, 'Chapter 1. Solutes and Solutions'), (70, '1.2. Topic'),
              (90, 'solubility increases with temperature.'), (300, 'Effect of temperature')]),
    ]
    page_range = {'start_page': 1, 'end_page': 2}
    text = rendered_text(doc, page_range, ['Solutes and Solutions', 'Effect of temperature'])
    assert norm('Temperature often affects how much solute a solvent can dissolve. solubility increases with temperature.') in text
    assert norm('Chapter 1. Solutes and Solutions') not in text
    assert norm('1.2. Topic') not in text
    assert norm('Effect of temperature') in text


def test_rendered_text_removes_decimal_section_header_between_split_lines():
    class Rect: height = 800
    class Page:
        rect = Rect()
        def __init__(self, lines): self.lines = lines
        def get_text(self, kind):
            return {'blocks': [{'lines': [
                {'bbox': (0, y, 100, y + 10), 'spans': [{'text': text}]}
                for y, text in self.lines]}]}
    doc = [Page([(700, 'The Little')]),
           Page([(40, '12.3. NIGHT SKY WATCHING'), (55, '49'), (90, "Dipper's handle points north.")])]
    text = rendered_text(doc, {'start_page': 1, 'end_page': 2}, [])
    assert norm("The Little Dipper's handle points north.") in text
    assert norm('12.3. NIGHT SKY WATCHING') not in text


def test_rendered_prose_validation_keeps_style_boundary_context():
    runs = [
        dict(kind='text', text='Clouds cause ', bold=False),
        dict(kind='text', text='precipitation', bold=True),
        dict(kind='text', text='-rain, snow or hail.', bold=False),
        dict(kind='math', text='x^2', bold=False),
        dict(kind='text', text='Further prose.', bold=False),
    ]
    assert prose_chunks(runs) == [
        'Clouds cause precipitation-rain, snow or hail.',
        'Further prose.',
    ]
