"""Focused Class 7 Science adapter and LyX-only regressions."""

import importlib.util
import os
import sys
from pathlib import Path

import pytest
from docx import Document
import hashlib
import json

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))


@pytest.fixture(scope="module")
def adapter():
    source_root = Path(os.environ.get("SCIENCE_INPUT_ROOT", ROOT))
    if not (source_root / "Class 7th Science").is_dir():
        pytest.skip("Class 7 Science supplied inputs are not available")
    spec = importlib.util.spec_from_file_location(
        "class7_science_practice_test", ROOT / "tools" / "rebuild_class7_science_practice.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_class7_adapter_maps_only_requested_chapters(adapter):
    assert adapter.CHAPTER_NUMBERS == (4, 5, 6, 7, 8, 9)
    assert adapter.MODULE == "Final-Science_Class7th"
    assert adapter.build.QUESTION_HEADINGS["Very Short Questions"] == "Very Short Answer Questions"
    for chapter in adapter.CHAPTER_NUMBERS:
        docx, pdf = adapter.question_paths(chapter)
        assert docx.name == f"Chapter-{chapter}-Questions-Class7th-Science.docx"
        assert pdf.name == f"Chapter-{chapter}-Questions-Class7th-Science.pdf"
        assert docx.is_file() and pdf.is_file()


def test_class7_source_hash_manifest_covers_notes_and_pairs(adapter):
    hashes = adapter.source_hashes()
    assert len(hashes) == 19
    assert hashes[adapter.WORKBOOK.name]
    for chapter in adapter.CHAPTER_NUMBERS:
        assert f"Chapter-{chapter}-Notes-Class7th-Science.pdf" in hashes
        assert f"Chapter-{chapter}-Questions-Class7th-Science.docx" in hashes
        assert f"Chapter-{chapter}-Questions-Class7th-Science.pdf" in hashes


def test_chapter8_omml_superscripts_and_fraction_operands_are_literal(adapter):
    source = adapter.QUESTIONS_DIR / "Chapter-8-Questions-Class7th-Science.docx"
    document = Document(source)
    values = [
        adapter.build.worksheet_paragraph_text(paragraph)
        for table in document.tables
        for row in table.rows
        for paragraph in row.cells[1].paragraphs
        if paragraph._element.xpath(".//m:oMath")
    ]
    assert values.count("d) m²") == 1
    assert values.count("d) m² /s") == 1
    assert any(r"\(\frac{\text{Totaldistancecovered}}{\text{Totaltimetaken}}\)" in value for value in values)


def test_all_class7_match_matrices_preserve_exact_native_cells(adapter, tmp_path):
    expected = {
        4: ["97042445cced3c939052e58877654c5052c801836a336c26cc5bc4c0324aac61", "647b4b1955b95b86e5c7b83c112fbb1c7b04f8a80f9b933dc26b98e6311f798f"],
        5: ["ff303e18aa4765b6f1fa027de13f75317061944b992c47a5c498284a2c88cadf", "e04f7d8275eccea13a80716b682451399a013fc2542cb25d33cd400fbc8c4ff3"],
        6: ["daa40d035a3b703115981601187f36b95e14d07cb14bb84027138c781cffc189", "b9ba114daf3ad881c8fdb0184d963f87acc914e47bd625fa57b2b04cb283f63f"],
        7: ["cac58c9a9b6c13de0adb56025ccc9dcb27ffc9062eb57b70a46d54b205155a7e", "cacb3fa5edbb48c14834afa830b4fb73be0a63c9a852e7e178ca4051aaea4048"],
        8: ["6c2bbd02c7be02ccb027957d6b77f4e292dd6fa96eacf9cc0b7d3a78cc5451d3", "14317ccb8baa633107ecb52b2990e4c1194b6f8c4ff4417aaa8a0e19e6156020"],
        9: ["b5c982812652401734b40ad34af3cb4d1509a857aa773b51b43388a111922a1f", "a9361d4058201a2ad8fb336a70741c52c619b0584601cbc70fb0f1ee5700d307"],
    }
    for chapter, hashes in expected.items():
        document = Document(adapter.QUESTIONS_DIR / f"Chapter-{chapter}-Questions-Class7th-Science.docx")
        actual = []
        for row in document.tables[1].rows:
            if row.cells[0].text.strip() not in {"35", "36"}:
                continue
            table = row.cells[1].tables[0]
            matrix = [[" ".join(cell.text.split()) for cell in native_row.cells] for native_row in table.rows]
            actual.append(hashlib.sha256(json.dumps(matrix, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest())
        assert actual == hashes


def test_class7_case_parts_strip_source_prefixes_but_audit_original(adapter):
    parsed, metadata = adapter.build.parse_case_events([
        ("text", "Passage."),
        ("text", "(a) First supplied part (1)"),
        ("text", "b. Second supplied part (2)"),
        ("text", "OR"),
    ])
    assert [part["text"] for part in parsed[-1][1]] == [
        "First supplied part", "Second supplied part"
    ]
    assert [item["source"] for item in metadata["removed_marks"]] == [
        "(a) First supplied part (1)", "b. Second supplied part (2)"
    ]


def test_lyx_only_compile_selection_skips_direct_tex(tmp_path, monkeypatch):
    import build_support

    calls = []

    def fake_command(args, directory, log, timeout=180, environment=None):
        calls.append(list(args))
        if args[0] == "lyx":
            (directory / args[-2]).write_text("\\documentclass{book}\\begin{document}\\end{document}")
        else:
            job = next(arg.split("=", 1)[1] for arg in args if arg.startswith("-jobname="))
            (directory / f"{job}.log").write_text("")
            import pymupdf
            pdf = pymupdf.open()
            pdf.new_page()
            pdf.save(directory / f"{job}.pdf")
            pdf.close()

    monkeypatch.setattr(build_support, "command", fake_command)
    outputs = build_support.compile_documents(
        tmp_path,
        "Final-Science_Class7th",
        {"lyx": "lyx", "pdflatex": "pdflatex"},
        formats=("lyx",),
    )
    assert set(outputs) == {"lyx"}
    assert not any("-jobname=Final-Science_Class7th-tex" in call for call in calls)
    assert any("-jobname=Final-Science_Class7th-lyx" in call for call in calls)
