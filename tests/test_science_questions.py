"""Focused regressions for semantic science question elements."""

import importlib.util
import hashlib
from pathlib import Path
import sys

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
ADAPTER = ROOT / "tools" / "rebuild_class8_science_practice.py"
sys.path.insert(0, str(ROOT / "tools"))

import convert
import science_table_manifest


@pytest.fixture(scope="module")
def adapter():
    import chapter_sources
    import scrape_chapters

    original_parse = chapter_sources.parse_chapter_sources
    original_download = scrape_chapters.download_content_images
    original_structure_check = scrape_chapters.check_structure
    spec = importlib.util.spec_from_file_location("science_practice_adapter", ADAPTER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    # Loading the production adapter installs build-time monkeypatches; restore
    # them so this focused parser test cannot affect other test modules.
    chapter_sources.parse_chapter_sources = original_parse
    scrape_chapters.download_content_images = original_download
    scrape_chapters.check_structure = original_structure_check
    return module


def test_case_event_parser_strips_question_marks_and_keeps_or_metadata(adapter):
    events = [
        ("text", "Passage about the case."),
        ("text", "First supplied part ( 1 )"),
        ("text", "Second supplied part (1)"),
        ("text", "Third supplied part (2)"),
        ("text", "OR"),
        ("text", "Alternative supplied part (2)"),
    ]

    parsed, metadata = adapter.parse_case_events(events)

    assert parsed == [
        ("text", "Passage about the case."),
        ("subquestions", [
            {"text": "First supplied part", "marks": 1, "answer_lines": 1},
            {"text": "Second supplied part", "marks": 1, "answer_lines": 1},
            {"text": "Third supplied part", "marks": 2, "answer_lines": 2},
            {"text": "Alternative supplied part", "marks": 2, "answer_lines": 2},
        ]),
    ]
    assert metadata == {
        "removed_marks": [
            {"source": "First supplied part ( 1 )", "marks": 1, "display": "First supplied part"},
            {"source": "Second supplied part (1)", "marks": 1, "display": "Second supplied part"},
            {"source": "Third supplied part (2)", "marks": 2, "display": "Third supplied part"},
            {"source": "Alternative supplied part (2)", "marks": 2, "display": "Alternative supplied part"},
        ],
        "omitted_or_count": 1,
        "case_part_count": 4,
    }


def test_unmarked_case_questions_keep_supplied_dengue_parts(adapter):
    events = [
        ("text", "Dengue case details."),
        ("text", "Answer the following:"),
        ("text", "Name the disease vector."),
        ("text", "State one prevention method."),
    ]

    parsed, metadata = adapter.parse_case_events(events)

    assert parsed == [
        ("text", "Dengue case details."),
        ("text", "Answer the following:"),
        ("subquestions", [
            {"text": "Name the disease vector.", "marks": None, "answer_lines": 1},
            {"text": "State one prevention method.", "marks": None, "answer_lines": 1},
        ]),
    ]
    assert metadata == {"removed_marks": [], "omitted_or_count": 0, "case_part_count": 2}


def test_actual_unmarked_dengue_four_part_question_from_class8_docx(adapter):
    source = ROOT / "Class 8th Science" / "Chapter-3-Questions-Class8th-Science.docx"
    document = Document(source)
    row = next(
        row
        for table in document.tables
        for row in table.rows
        if row.cells and row.cells[0].text.strip() == "57"
    )
    events = [
        ("text", adapter.clean(line))
        for line in row.cells[1].text.splitlines()
        if adapter.clean(line)
    ]

    parsed, metadata = adapter.parse_case_events(events)

    assert parsed[:-1] == [
        ("text", "The following table contains information about the number of dengue cases reported in a hospital over a period of one year:"),
        ("text", "Make a bar graph of the number of cases on the Y-axis and the month on the X-axis. Critically analyse your findings and answer the following:"),
    ]
    parts = parsed[-1]
    assert parts[0] == "subquestions"
    assert [item["text"] for item in parts[1]] == [
        "In which three months were the dengue cases highest?",
        "In which month(s) were the cases lowest?",
        "What natural or environmental factors during the peak months might contribute to the increase in dengue cases?",
        "Suggest a few preventive steps that the community or government can take before the peak season to reduce the spread of dengue.",
    ]
    assert all(item["marks"] is None and item["answer_lines"] == 1 for item in parts[1])
    assert metadata == {"removed_marks": [], "omitted_or_count": 0, "case_part_count": 4}


def test_class8_chapter9_omml_superscripts_are_retained_in_all_three_prompts(adapter, tmp_path):
    source = ROOT / "Class 8th Science" / "Chapter-9-Questions-Class8th-Science.docx"
    document = Document(source)
    prompts = {
        row.cells[0].text.strip(): adapter.row_question(row)[1]
        for table in document.tables
        for row in table.rows
        if row.cells[0].text.strip() in {"28", "30", "32"}
    }

    assert prompts == {
        "28": ["An object has a mass of 400 g and a volume of 40 cm³ . What is its density?"],
        "30": ["A stone sculpture weighs 225 g and has a volume of 90 cm³ . Calculate its density and predict whether it will float or sink in water."],
        "32": ["Calculate the density of a stone with mass 50 g and volume 20 cm³ ."],
    }

    question_elements = [
        {"type": "question", "content": [{"type": "body", "segments": [(line, False)]}]}
        for lines in prompts.values()
        for line in lines
    ]
    tex_path = tmp_path / "omml.tex"
    lyx_path = tmp_path / "omml.lyx"
    convert.write_tex(question_elements, tex_path, tmp_path)
    convert.write_lyx(question_elements, lyx_path, tmp_path)

    for output in (tex_path, lyx_path):
        rendered = output.read_text(encoding="utf-8")
        assert rendered.count("cm³") == 3


def test_class8_chapter12_actual_omml_arrow_is_retained(adapter):
    source = ROOT / "Class 8th Science" / "Chapter-12-Questions-Class8th-Science.docx"
    document = Document(source)
    paragraphs = [
        paragraph
        for table in document.tables
        for row in table.rows
        for paragraph in row.cells[1].paragraphs
        if paragraph._element.xpath(".//m:oMath")
    ]

    extracted = [adapter.worksheet_paragraph_text(paragraph) for paragraph in paragraphs]

    assert any("grass → grasshopper → frog → snake → eagle" in text for text in extracted)
    assert all("\u200b" not in text for text in extracted)


def test_unsupported_worksheet_omml_fails_closed(adapter):
    from docx.oxml import OxmlElement

    paragraph = Document().add_paragraph("Unknown: ")
    paragraph._p.append(OxmlElement("m:f"))

    with pytest.raises(ValueError, match="Unsupported worksheet OMML element: f"):
        adapter.worksheet_paragraph_text(paragraph)


def test_semantic_question_keeps_mixed_content_and_answer_line_count():
    content = [
        {"type": "body", "segments": [("Explain the diagram.", False)]},
        {"type": "image", "filename": "question.png", "scale": 0.25},
        {"type": "table", "rows": [["Part", "Value"], ["A", "1"]]},
        {"type": "answer_lines", "count": 2},
        {"type": "subquestions", "items": [
            {"content": [
                {"type": "body", "segments": [("Name the part.", False)]},
                {"type": "answer_lines", "count": 1},
            ]},
        ]},
    ]
    question = {"type": "question", "content": content}

    assert question["type"] == "question"
    assert [item["type"] for item in question["content"]] == [
        "body", "image", "table", "answer_lines", "subquestions"
    ]
    assert question["content"][3]["count"] == 2
    assert question["content"][4]["items"][0]["content"][1]["count"] == 1


def test_reviewed_science_table_requires_known_source_hash():
    image = {
        "filename": "ch8_questions_rId26.png",
        "sha256": "7bca87777c3031804e3c4c4aa643c22cd56bcbd4a675f73c25cf46191bb8091b",
    }

    table = science_table_manifest.convert_table(image, 8, 3)

    assert table["type"] == "table"
    assert table["source_identity"] == "8:3:rId26"
    assert table["rows"][0] == ["S. No.", "1", "2", "3", "4", "5", "6"]
    with pytest.raises(ValueError, match="Reviewed table image changed"):
        science_table_manifest.convert_table({**image, "sha256": "0" * 64}, 8, 3)


def test_actual_class8_table_keeps_original_media_hash_but_renders_native_table(adapter, tmp_path):
    source = ROOT / "Class 8th Science" / "Chapter-3-Questions-Class8th-Science.docx"
    document = Document(source)
    media_dir = tmp_path / "media"
    source_images = [
        image
        for row in document.tables[1].rows
        for paragraph in row.cells[1].paragraphs
        for image in adapter.paragraph_images(document, paragraph, media_dir, 3)
        if image["filename"].endswith("rId20.png")
    ]

    assert len(source_images) == 1
    image = source_images[0]
    original = media_dir / image["original_filename"]
    assert hashlib.sha256(original.read_bytes()).hexdigest() == image["sha256"]
    table = science_table_manifest.convert_table(image, 8, 3)
    assert table["type"] == "table"

    elements = [{"type": "question", "content": [
        {"type": "body", "segments": [("Match the following:", False)]},
        table,
    ]}]
    tex_path = tmp_path / "native-table.tex"
    lyx_path = tmp_path / "native-table.lyx"
    convert.write_tex(elements, tex_path, tmp_path)
    convert.write_lyx(elements, lyx_path, tmp_path)

    for output in (tex_path, lyx_path):
        rendered = output.read_text(encoding="utf-8")
        assert image["filename"] not in rendered
        assert "Match the following:" in rendered
        assert "Column A" in rendered


def test_tex_writer_preserves_question_images_and_tables(tmp_path):
    output = tmp_path / "mixed-question.tex"
    convert.write_tex([{"type": "question", "content": [
        {"type": "body", "segments": [("Read the figure and table.", False)]},
        {"type": "image", "filename": "case.png", "scale": 0.25},
        {"type": "table", "rows": [["Part", "Value"], ["A", "1"]]},
    ]}], output, tmp_path)

    rendered = output.read_text(encoding="utf-8")
    assert "media/case.png" in rendered
    assert "Read the figure and table." in rendered
    assert "Part" in rendered and "Value" in rendered and "A" in rendered


def test_tex_question_writer_uses_unbolded_body_and_linewidth_rules(tmp_path):
    output = tmp_path / "questions.tex"
    convert.write_tex([
        {"type": "question", "content": [
            {"type": "body", "segments": [("Question text", False)]},
            {"type": "answer_lines", "count": 2},
        ]},
    ], output, tmp_path)

    rendered = output.read_text(encoding="utf-8")

    assert r"\textbf{Question text}" not in rendered
    assert "Question text" in rendered
    assert rendered.count(r"\rule{\linewidth}{0.3pt}") == 2


def test_tex_semantic_question_writer_bounds_running_headers(tmp_path):
    output = tmp_path / "bounded-headers.tex"
    convert.write_tex([{"type": "question", "content": [
        {"type": "body", "segments": [("Question text", False)]},
    ]}], output, tmp_path)

    rendered = output.read_text(encoding="utf-8")

    assert r"\fancyhead[LE,RO]{\thepage}" in rendered
    assert r"\fancyhead[LO]{\parbox[b]{.88\textwidth}{\raggedright\small\rightmark}}" in rendered
    assert r"\fancyhead[RE]{\parbox[b]{.88\textwidth}{\raggedleft\small\leftmark}}" in rendered
    assert r"\setlength{\headheight}{32pt}" in rendered


def test_copied_lyx_template_bounds_running_headers(tmp_path):
    convert.copy_template_assets(convert.TEMPLATE_DIR, tmp_path)

    structure = (tmp_path / "structure.tex").read_text(encoding="utf-8")

    assert r"\fancyhead[LO]{\parbox[b]{.88\textwidth}{\raggedright\rightmark}}" in structure
    assert r"\fancyhead[RE]{\parbox[b]{.88\textwidth}{\raggedleft\leftmark}}" in structure
    assert r"\setlength{\headheight}{32pt}" in structure


def test_tex_writer_ragged_right_applies_only_to_fillblank_body_and_keeps_blanks(tmp_path):
    output = tmp_path / "fill-blanks.tex"
    convert.write_tex([
        {"type": "question", "content": [
            {"type": "body", "segments": [("Fill: ____ and ______.", False)], "ragged_right": True},
        ]},
        {"type": "question", "content": [
            {"type": "body", "segments": [("Ordinary body.", False)]},
        ]},
    ], output, tmp_path)

    rendered = output.read_text(encoding="utf-8")

    assert rendered.count(r"{\raggedright ") == 1
    assert r"{\raggedright Fill: \_\_\_\_ and \_\_\_\_\_\_.\par}" in rendered
    assert "Ordinary body.\\par" in rendered


def test_lyx_writer_uses_native_nested_enumerate_for_subquestions(tmp_path):
    output = tmp_path / "questions.lyx"
    convert.write_lyx([{"type": "question", "content": [
        {"type": "body", "segments": [("Case prompt", False)]},
        {"type": "subquestions", "items": [
            {"content": [{"type": "body", "segments": [("First part", False)]}]},
            {"content": [{"type": "body", "segments": [("Second part", False)]}]},
        ]},
    ]}], output, tmp_path)

    rendered = output.read_text(encoding="utf-8")

    assert rendered.count("\\begin_layout Enumerate") == 3
    assert "First part" in rendered and "Second part" in rendered
    assert "\\begin_inset ERT" not in rendered


def test_lyx_writer_preserves_reviewed_merged_header_as_native_table(tmp_path):
    output = tmp_path / "merged-header.lyx"
    convert.write_lyx([{"type": "question", "content": [
        {"type": "body", "segments": [("Read the temperature table.", False)]},
        {"type": "table", "rows": [
            ["Temperature at", "", "", "", "", "", ""],
            ["DAY", "7 am", "10 am", "1 pm", "4 pm", "7 pm", "10 pm"],
        ], "merged_header": True},
    ]}], output, tmp_path)

    rendered = output.read_text(encoding="utf-8")

    assert '<lyxtabular version="3" rows="2" columns="7">' in rendered
    assert rendered.count("<row>") == 2
    assert rendered.count("<cell") == 14
    assert 'multicolumn="1"' in rendered
    assert 'multicolumn="2"' in rendered
    assert "Temperature at" in rendered
    assert "10 pm" in rendered
