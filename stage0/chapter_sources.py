"""Structured chapter inputs and infographic-PDF helpers."""
from html import unescape
from pathlib import Path
import re
from urllib.parse import urlparse

from PIL import Image, ImageChops, ImageStat


WORKBOOK_HEADERS = (
    'Chapter Name',
    'Chapter-Number',
    'Important Points and Formulas link',
    'Infographic link',
)
INFOGRAPHIC_PDF_RE = re.compile(
    r'https:\\?/\\?/cn[.]edurev[.]in/files/[^"\'\\\s<>]+[.]pdf(?:[?][^"\'\\\s<>]*)?',
    re.I,
)


def normalize_edurev_url(value, expected_path):
    """Return a validated EduRev URL, repairing only a missing scheme."""
    text = str(value or '').strip()
    if re.match(r'^edurev[.]in/', text, re.I):
        text = 'https://' + text
    parsed = urlparse(text)
    if parsed.scheme not in ('http', 'https') or (parsed.hostname or '').lower() not in ('edurev.in', 'www.edurev.in'):
        raise ValueError(f'Expected an HTTP(S) EduRev URL, got: {text or "<blank>"}')
    if not parsed.path.startswith(expected_path):
        raise ValueError(f'Expected an EduRev {expected_path} URL, got: {text}')
    return text


def parse_workbook(path, sheet):
    """Read formula links and optional infographics from an explicit sheet."""
    if not sheet:
        raise ValueError('--sheet is required when --links is an .xlsx workbook')
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError('Workbook input requires openpyxl; install requirements.txt') from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet not in workbook.sheetnames:
        raise ValueError(f'Workbook sheet not found: {sheet}')
    worksheet = workbook[sheet]
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = tuple(str(value or '').strip() for value in next(rows))
    except StopIteration as exc:
        raise ValueError(f'Workbook sheet is empty: {sheet}') from exc
    # Some supplied sheets repeat the chapter name in a second display column.
    # Recognize that exact layout without shifting unrelated/malformed schemas.
    duplicate_name = headers == ('Chapter Name',) + WORKBOOK_HEADERS
    if headers != WORKBOOK_HEADERS and not duplicate_name:
        raise ValueError('Workbook headers must be: ' + ', '.join(WORKBOOK_HEADERS))
    chapters = []
    for row_number, row in enumerate(rows, start=2):
        offset = 1 if duplicate_name else 0
        selected = row[offset:offset + 4]
        values = tuple(selected) + (None,) * max(0, 4 - len(selected))
        if not any(str(value or '').strip() for value in values):
            continue
        if any(not str(value or '').strip() for value in values[:3]):
            raise ValueError(f'Workbook row {row_number} has a missing required value')
        name, chapter_label, formula, infographic = (str(value or '').strip() for value in values)
        if duplicate_name and str(row[0] or '').strip():
            # The first name column contains the supplied display title; the
            # second may contain spelling/capitalization variants.
            name = str(row[0]).strip()
        match = re.fullmatch(r'Chapter-(\d+)', chapter_label, re.I)
        if not match:
            raise ValueError(f'Workbook row {row_number} has an invalid chapter number: {chapter_label}')
        if name.casefold() == 'periimeter and area':
            name = 'Perimeter and Area'
        chapters.append({
            'num': int(match.group(1)),
            'name': name,
            'url': normalize_edurev_url(formula, '/t/'),
            'formula_url': normalize_edurev_url(formula, '/t/'),
            'infographic_url': normalize_edurev_url(infographic, '/p/') if infographic else None,
        })
    if not chapters:
        raise ValueError(f'Workbook sheet has no chapter rows: {sheet}')
    return chapters


def parse_chapter_sources(path, sheet=None):
    """Parse legacy text links or paired links from an XLSX workbook."""
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        return parse_workbook(path, sheet)
    if sheet:
        raise ValueError('--sheet can only be used with an .xlsx workbook')
    from scrape_images import parse_links
    chapters = parse_links(path)
    for chapter in chapters:
        chapter['formula_url'] = chapter['url']
        chapter['infographic_url'] = None
    return chapters


def extract_infographic_pdf_url(html):
    """Extract exactly one approved EduRev file PDF from an infographic page."""
    normalized = unescape(html).replace(r'\/', '/')
    urls = sorted(set(INFOGRAPHIC_PDF_RE.findall(normalized)))
    if len(urls) != 1:
        raise ValueError(f'Expected exactly one embedded infographic PDF, found {len(urls)}')
    parsed = urlparse(urls[0])
    if parsed.scheme != 'https' or parsed.hostname != 'cn.edurev.in' or not parsed.path.startswith('/files/'):
        raise ValueError('Infographic PDF URL is outside the approved EduRev file host')
    return urls[0]


def validate_pdf(data):
    """Reject empty, corrupt, password-protected, or dimensionless PDFs."""
    import pymupdf
    with pymupdf.open(stream=data, filetype='pdf') as document:
        if document.needs_pass or not len(document):
            raise ValueError('Infographic PDF is empty or password protected')
        if any(page.rect.width <= 0 or page.rect.height <= 0 for page in document):
            raise ValueError('Infographic PDF contains an invalid page')


def render_pdf_pages(data, dpi=200):
    """Render every PDF page to an ordered, opaque PNG."""
    import pymupdf
    scale = dpi / 72
    rendered = []
    with pymupdf.open(stream=data, filetype='pdf') as document:
        for page in document:
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
            rendered.append(pixmap.tobytes('png'))
    return rendered


def find_reference_pdf(directory, chapter_number, kind):
    """Resolve one chapter reference despite punctuation/case differences."""
    directory = Path(directory)
    matches = []
    for path in directory.glob('*.pdf'):
        key = re.sub(r'[^a-z0-9]', '', path.name.lower())
        if f'chapter{chapter_number}' not in key:
            continue
        if kind == 'formula' and 'important' in key and 'formula' in key:
            matches.append(path)
        if kind == 'infographic' and 'infographic' in key:
            matches.append(path)
    if len(matches) != 1:
        raise ValueError(f'Expected one Chapter {chapter_number} {kind} reference PDF, found {len(matches)}')
    return matches[0]


def compare_pdf_pages(source_data, reference_path, rms_limit=15.0):
    """Compare rendered page appearance after normalizing sub-point dimensions."""
    import pymupdf

    def grayscale(document, index):
        pixmap = document[index].get_pixmap(colorspace=pymupdf.csGRAY, alpha=False)
        return Image.frombytes('L', (pixmap.width, pixmap.height), pixmap.samples)

    with pymupdf.open(stream=source_data, filetype='pdf') as source, pymupdf.open(reference_path) as reference:
        if len(source) != len(reference):
            raise ValueError(f'Infographic page count differs: source {len(source)}, reference {len(reference)}')
        values = []
        for index in range(len(source)):
            source_image = grayscale(source, index)
            reference_image = grayscale(reference, index).resize(source_image.size, Image.Resampling.BICUBIC)
            rms = ImageStat.Stat(ImageChops.difference(source_image, reference_image)).rms[0]
            values.append(round(rms, 2))
        if any(value > rms_limit for value in values):
            raise ValueError(f'Infographic visual difference exceeds RMS {rms_limit}: {values}')
        return values
